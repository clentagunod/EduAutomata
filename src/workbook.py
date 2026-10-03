import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .constants import CATEGORIES
from .storage import SQLiteRepository


def excel_value(value: str | None):
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    return int(number) if number.is_integer() else number


class WorkbookSynchronizer:
    def sync(self, repository: SQLiteRepository, workbook_path: Path) -> Path:
        workbook_path = workbook_path.expanduser()
        workbook_path.parent.mkdir(parents=True, exist_ok=True)

        if workbook_path.exists():
            workbook = load_workbook(workbook_path)
        else:
            workbook = Workbook()
            if workbook.sheetnames == ["Sheet"]:
                del workbook["Sheet"]

        sheet_names = (*CATEGORIES, "Learners")
        for category in sheet_names:
            if category in workbook.sheetnames:
                del workbook[category]

        exports = []
        for category in CATEGORIES:
            dataset = repository.load(category)
            exports.append((category, dataset.headers, dataset.rows))
        learner_headers, learner_rows = repository.load_learners()
        exports.append(("Learners", learner_headers, learner_rows))
        for index, (category, headers, records) in enumerate(exports):
            worksheet = workbook.create_sheet(category, index)
            worksheet.sheet_view.showGridLines = False
            worksheet.append(headers)
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(vertical="center")
            for row in records:
                worksheet.append([excel_value(row.get(header)) for header in headers])
            worksheet.freeze_panes = "B2"
            if worksheet.max_row and worksheet.max_column:
                worksheet.auto_filter.ref = worksheet.dimensions
                for column_index, cells in enumerate(worksheet.columns, start=1):
                    width = min(
                        36,
                        max(12, max(len(str(cell.value or "")) for cell in cells) + 2),
                    )
                    worksheet.column_dimensions[
                        get_column_letter(column_index)
                    ].width = width

        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=workbook_path.parent,
                suffix=".xlsx",
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            os.replace(temporary_path, workbook_path)
        finally:
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()
            workbook.close()
        return workbook_path
