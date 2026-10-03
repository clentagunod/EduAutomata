import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from source.analytics import summarize
from source.constants import CATEGORIES, DEFAULT_WORKBOOK_PATH
from source.data import category_columns
from source.preferences import Preferences, PreferencesStore
from source.storage import SQLiteRepository
from source.workbook import WorkbookSynchronizer


class DatabaseAnalyticsTests(unittest.TestCase):
    def test_category_columns_support_different_source_headers(self):
        self.assertEqual(
            category_columns("Attendance", ["Learners", "Week 2", "Week 1"]),
            ["Week 1", "Week 2"],
        )
        self.assertEqual(
            category_columns("Activity", ["Learners", "Activity 2", "Activity 1"]),
            ["Activity 1", "Activity 2"],
        )
        self.assertEqual(
            category_columns("Assessment", ["Learners", "Assessment2", "Assessment1"]),
            ["Assessment1", "Assessment2"],
        )

    def test_attendance_summary_computes_week_rates_and_totals(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "Attendance.csv"
            path.write_text(
                "Learners,Week 1,Week 2\nA,1,0\nB,1,1\nC,0,1\n",
                encoding="utf-8",
            )
            repository = SQLiteRepository(root / "records.sqlite3")
            repository.import_legacy_csvs({"Attendance": path}, root / "missing.csv")
            dataset = repository.load("Attendance")
            summary = summarize(dataset)

        self.assertEqual(summary.learner_count, 3)
        self.assertAlmostEqual(summary.attendance_rate, 100 * 4 / 6)
        self.assertEqual((summary.weeks[0].present, summary.weeks[0].absent), (2, 1))
        self.assertEqual(summary.weeks[1].average, 2 / 3)

    def test_score_summary_ignores_empty_and_non_numeric_values(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "Activity.csv"
            path.write_text(
                "Learners,Activity 1,Activity 2\nA,5,\nB,3,not graded\nC,,4\n",
                encoding="utf-8",
            )
            repository = SQLiteRepository(root / "records.sqlite3")
            repository.import_legacy_csvs({"Activity": path}, root / "missing.csv")
            summary = summarize(repository.load("Activity"))

        self.assertEqual(summary.values, [5.0, 3.0, 4.0])
        self.assertEqual(summary.average, 4.0)
        self.assertEqual(summary.median, 4.0)
        self.assertEqual(summary.weeks[1].count, 1)

    def test_add_learner_updates_metadata_and_all_record_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            repository = SQLiteRepository(Path(directory) / "records.sqlite3")
            repository.add_learner("New Student", "M", "9", "Avocado")

            self.assertEqual(
                repository.load_learners()[1][0],
                {
                    "Learners": "New Student",
                    "Sex": "M",
                    "Grade-Section": "9 - Avocado",
                },
            )
            for category in CATEGORIES:
                dataset = repository.load(category)
                self.assertEqual(dataset.rows[0]["Learners"], "New Student")
                self.assertTrue(
                    all(
                        dataset.rows[0][period] == "" for period in dataset.week_columns
                    )
                )

    def test_add_learner_rejects_duplicate_names(self):
        with tempfile.TemporaryDirectory() as directory:
            repository = SQLiteRepository(Path(directory) / "records.sqlite3")
            repository.add_learner("Existing", "F", "8", "Rafflesia")
            with self.assertRaisesRegex(ValueError, "already exists"):
                repository.add_learner(" existing ", "F", "9", "Avocado")

    def test_delete_learner_removes_rows_and_synced_workbook_entries(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            repository = SQLiteRepository(root / "records.sqlite3")
            repository.add_learner("Existing", "F", "8", "Rafflesia")
            repository.add_learner("To Delete", "M", "9", "Avocado")
            for category in CATEGORIES:
                dataset = repository.load(category)
                dataset.rows[1][dataset.week_columns[0]] = "2"
                repository.save(dataset)
            repository.delete_learner(" to delete ")
            self.assertEqual(
                [row["Learners"] for row in repository.load_learners()[1]],
                ["Existing"],
            )
            workbook_path = root / "combined.xlsx"
            WorkbookSynchronizer().sync(repository, workbook_path)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                for category in (*CATEGORIES, "Learners"):
                    worksheet = workbook[category]
                    self.assertEqual(worksheet.max_row, 2)
                    self.assertEqual(worksheet["A2"].value, "Existing")
            finally:
                workbook.close()


class SQLiteRepositoryTests(unittest.TestCase):
    def test_fresh_database_has_blank_periods_and_no_learners(self):
        with tempfile.TemporaryDirectory() as directory:
            repository = SQLiteRepository(Path(directory) / "records.sqlite3")

            self.assertEqual(repository.load_learners()[1], [])
            for category in CATEGORIES:
                dataset = repository.load(category)
                self.assertEqual(len(dataset.week_columns), 8)
                self.assertEqual(dataset.rows, [])

    def test_records_persist_and_learner_deletion_cascades(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "data" / "records.sqlite3"
            repository = SQLiteRepository(database_path)
            with repository.connection() as connection:
                tables = {
                    row["name"]
                    for row in connection.execute(
                        "SELECT name FROM sqlite_master WHERE type = 'table'"
                    )
                }
            self.assertTrue(
                {
                    "learners",
                    "attendance_records",
                    "activity_records",
                    "assessment_records",
                }.issubset(tables)
            )
            repository.add_learner("Learner A", "F", "8", "Rafflesia")
            dataset = repository.load("Attendance")
            dataset.rows[0]["Week 1"] = "1"
            repository.save(dataset)

            reloaded = SQLiteRepository(database_path)
            self.assertEqual(reloaded.load("Attendance").rows[0]["Week 1"], "1")
            reloaded.delete_learner("learner a")
            self.assertEqual(reloaded.load("Attendance").rows, [])
            self.assertEqual(reloaded.load_learners()[1], [])

    def test_legacy_csvs_import_once_and_preserve_periods_and_scores(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            learners_path = root / "Learners.csv"
            learners_path.write_text(
                "Learners,Sex,Grade-Section\nLearner A,F,8 - Rafflesia\n",
                encoding="utf-8",
            )
            csv_paths = {}
            for category, header in (
                ("Attendance", "Week 2"),
                ("Activity", "Activity 3"),
                ("Assessment", "Assessment1"),
            ):
                path = root / f"{category}.csv"
                path.write_text(f"Learners,{header}\nLearner A,4\n", encoding="utf-8")
                csv_paths[category] = path

            repository = SQLiteRepository(root / "records.sqlite3")
            self.assertTrue(repository.import_legacy_csvs(csv_paths, learners_path))
            self.assertFalse(repository.import_legacy_csvs(csv_paths, learners_path))
            self.assertEqual(repository.load("Attendance").week_columns, ["Week 2"])
            self.assertEqual(repository.load("Activity").rows[0]["Activity 3"], "4")
            self.assertEqual(
                repository.load_learners()[1][0]["Grade-Section"], "8 - Rafflesia"
            )


class PersistenceTests(unittest.TestCase):
    def test_preferences_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "preferences.json"
            preferences = Preferences.defaults()
            preferences.workbook_path = Path(directory) / "combined.xlsx"
            preferences.dark_mode = True
            preferences.category = "Assessment"
            preferences.window_geometry = "0102"
            store = PreferencesStore(path)
            store.save(preferences)
            loaded = store.load()
            saved_preferences = json.loads(path.read_text(encoding="utf-8"))

        self.assertFalse({"csv_paths", "learners_path"}.intersection(saved_preferences))
        self.assertEqual(loaded.workbook_path, preferences.workbook_path)
        self.assertTrue(loaded.dark_mode)
        self.assertEqual(loaded.category, "Assessment")
        self.assertEqual(loaded.window_geometry, "0102")

    def test_legacy_csv_paths_are_read_only_migration_settings(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "preferences.json"
            custom_csv_paths = {
                category: str(root / f"custom-{category}.csv")
                for category in CATEGORIES
            }
            path.write_text(
                json.dumps(
                    {
                        "csv_paths": custom_csv_paths,
                        "learners_path": str(root / "custom-learners.csv"),
                    }
                ),
                encoding="utf-8",
            )
            preferences = PreferencesStore(path).load()

        self.assertEqual(
            preferences.legacy_csv_paths,
            {category: Path(value) for category, value in custom_csv_paths.items()},
        )
        self.assertEqual(preferences.legacy_learners_path, root / "custom-learners.csv")

    def test_legacy_workbook_default_migrates_to_synced_folder(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "preferences.json"
            path.write_text(
                '{"workbook_path": "'
                + str(DEFAULT_WORKBOOK_PATH.parents[1] / "ARAL_Synced.xlsx")
                + '"}',
                encoding="utf-8",
            )
            loaded = PreferencesStore(path).load()

        self.assertEqual(loaded.workbook_path, DEFAULT_WORKBOOK_PATH)

    def test_workbook_has_a_sheet_for_every_category(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            repository = SQLiteRepository(root / "records.sqlite3")
            repository.add_learner("Learner A", "F", "8", "Rafflesia")
            for category in CATEGORIES:
                dataset = repository.load(category)
                dataset.rows[0][dataset.week_columns[0]] = "5"
                repository.save(dataset)
            workbook_path = root / "combined.xlsx"
            WorkbookSynchronizer().sync(repository, workbook_path)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                self.assertTrue(
                    set((*CATEGORIES, "Learners")).issubset(workbook.sheetnames)
                )
                self.assertEqual(workbook["Activity"]["B2"].value, 5)
                self.assertEqual(workbook["Assessment"]["A2"].value, "Learner A")
                self.assertEqual(workbook["Learners"]["A2"].value, "Learner A")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
